"""Valid document fixtures generated with format-aware libraries, not renamed ZIPs."""
import hashlib
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile

from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from odf.opendocument import OpenDocumentText, load
from odf.text import P

from lynowcut.engine import Options, process


class DocumentRoundtripTests(unittest.TestCase):
    def test_real_documents_reopen_and_keep_content(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / 'photo.jpg'
            Image.effect_noise((800, 600), 70).convert('RGB').save(image, quality=99)
            doc = Document(); doc.add_paragraph('LynowCut roundtrip'); doc.add_picture(str(image))
            doc.save(root / 'document.docx')
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = 'LynowCut roundtrip'
            slide.shapes.add_picture(str(image), 0, Inches(1), width=Inches(5)); deck.save(root / 'slides.pptx')
            book = Workbook(); book.active['A1'] = 'LynowCut roundtrip'
            book.active.add_image(SheetImage(str(image)), 'A3'); book.save(root / 'sheet.xlsx')
            odt = OpenDocumentText(); odt.text.addElement(P(text='LynowCut roundtrip'))
            odt.save(str(root / 'document.odt'))
            sources = [root / name for name in ('document.docx', 'slides.pptx', 'sheet.xlsx', 'document.odt')]
            for src in sources:
                with self.subTest(format=src.suffix):
                    digest = hashlib.sha256(src.read_bytes()).digest()
                    result = process(src, root / 'out', Options())
                    output = Path(result.output) if result.output else src
                    self.assertEqual(hashlib.sha256(src.read_bytes()).digest(), digest)
                    with zipfile.ZipFile(src) as original, zipfile.ZipFile(output) as converted:
                        self.assertIsNone(converted.testzip())
                        self.assertEqual(original.namelist(), converted.namelist())
                        for name in original.namelist():
                            if name.endswith(('.xml', '.rels')):
                                self.assertEqual(original.read(name), converted.read(name))
                    if src.suffix == '.docx': self.assertEqual(Document(output).paragraphs[0].text, 'LynowCut roundtrip')
                    elif src.suffix == '.pptx': self.assertEqual(Presentation(output).slides[0].shapes[0].text, 'LynowCut roundtrip')
                    elif src.suffix == '.xlsx': self.assertEqual(load_workbook(output).active['A1'].value, 'LynowCut roundtrip')
                    else: self.assertTrue(load(str(output)).text.childNodes)
                    # Release CI requires an independent office application as well.
                    if os.environ.get('LIBREOFFICE'):
                        destination = root / ('render-' + src.suffix[1:]); destination.mkdir()
                        subprocess.run([os.environ['LIBREOFFICE'], '-env:UserInstallation=' + (root / ('profile-' + src.suffix[1:])).as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(destination), str(output)], check=True, timeout=120, capture_output=True)
                        self.assertTrue(any(p.stat().st_size > 0 for p in destination.glob('*.pdf')), 'LibreOffice could not render document')
